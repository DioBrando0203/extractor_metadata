"""Pruebas sintéticas para extractores de adjuntos; no contienen correos reales."""

from __future__ import annotations

import io
import subprocess
import threading
import time
import zipfile

import ezdxf
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from app.core.config import settings
from app.services import metadata
from app.services.metadata import exiftool, extractor, results


def _without_exiftool(monkeypatch) -> None:
    monkeypatch.setattr(
        extractor,
        "exiftool_metadata",
        lambda _payload, _filename, *, deadline=None: results.ExtractionResult(),
    )


def _values(items: list) -> dict[str, str]:
    return {item.label: item.value for item in items}


class _FakeProcess:
    def __init__(self, stdout: object, *, hangs: bool = False) -> None:
        self.stdin = io.BytesIO()
        self.stdout = stdout
        self.returncode: int | None = None
        self.hangs = hangs
        self.killed = False

    def poll(self) -> int | None:
        return self.returncode

    def wait(self, timeout: float | None = None) -> int:
        if self.returncode is None and self.hangs:
            raise subprocess.TimeoutExpired("exiftool", timeout)
        if self.returncode is None:
            self.returncode = 0
        return self.returncode

    def kill(self) -> None:
        self.killed = True
        self.returncode = -9

    def terminate(self) -> None:
        self.kill()


class _BlockingStdout:
    def __init__(self) -> None:
        self.closed = threading.Event()

    def read(self, _size: int) -> bytes:
        self.closed.wait(timeout=2)
        return b""

    def close(self) -> None:
        self.closed.set()


def test_extracts_image_dimensions_and_exif(monkeypatch) -> None:
    _without_exiftool(monkeypatch)
    image = Image.new("RGB", (7, 5), color="navy")
    exif = Image.Exif()
    exif[315] = "Equipo de calidad"
    output = io.BytesIO()
    image.save(output, format="JPEG", exif=exif)

    content_type, items, warnings = metadata.extract_file_metadata(output.getvalue(), "foto.jpg")

    values = _values(items)
    assert content_type == "image/jpeg"
    assert values["Dimensiones"] == "7 × 5"
    assert values["Artist"] == "Equipo de calidad"
    assert not warnings


def test_extracts_pdf_properties_and_uses_signature_over_extension(monkeypatch) -> None:
    _without_exiftool(monkeypatch)
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.add_metadata({"/Title": "Informe técnico", "/Author": "QA"})
    output = io.BytesIO()
    writer.write(output)

    content_type, items, warnings = metadata.extract_file_metadata(
        output.getvalue(), "aparenta-ser-imagen.jpg"
    )

    values = _values(items)
    assert content_type == "application/pdf"
    assert values["Páginas"] == "1"
    assert values["Title"] == "Informe técnico"
    assert not warnings


def test_extracts_xlsx_properties_without_reading_worksheets(monkeypatch) -> None:
    _without_exiftool(monkeypatch)
    book = Workbook()
    book.active.title = "Resumen"
    book.create_sheet("Datos")
    book.properties.creator = "Analista"
    book.properties.title = "Matriz"
    output = io.BytesIO()
    book.save(output)
    book.close()

    content_type, items, warnings = metadata.extract_file_metadata(output.getvalue(), "matriz.xlsx")

    values = _values(items)
    assert content_type and "spreadsheetml" in content_type
    assert values["Cantidad de hojas"] == "2"
    assert values["Hojas"] == "Resumen, Datos"
    assert values["Autor"] == "Analista"
    assert not warnings


def test_extracts_docx_core_properties(monkeypatch) -> None:
    _without_exiftool(monkeypatch)
    document = Document()
    document.core_properties.author = "Redacción"
    document.core_properties.title = "Acta"
    document.core_properties.keywords = "obra, inspección"
    output = io.BytesIO()
    document.save(output)

    content_type, items, warnings = metadata.extract_file_metadata(output.getvalue(), "acta.docx")

    values = _values(items)
    assert content_type and "wordprocessingml" in content_type
    assert values["Formato"] == "Word Open XML"
    assert values["Autor"] == "Redacción"
    assert values["Título"] == "Acta"
    assert values["keywords"] == "obra, inspección"
    assert not warnings


def test_extracts_dxf_header_and_layers(monkeypatch) -> None:
    _without_exiftool(monkeypatch)
    document = ezdxf.new("R2010")
    document.header["$INSUNITS"] = 4
    document.layers.new("REVISION")
    output = io.StringIO()
    document.write(output)

    content_type, items, warnings = metadata.extract_file_metadata(
        output.getvalue().encode("latin-1"), "plano.dxf"
    )

    values = _values(items)
    assert content_type == "application/dxf"
    assert values["Formato"] == "AC1024"
    assert values["Unidades"] == "4"
    assert int(values["Capas"]) >= 2
    assert not warnings


def test_corrupt_pdf_returns_generic_metadata_and_warning(monkeypatch) -> None:
    _without_exiftool(monkeypatch)

    content_type, items, warnings = metadata.extract_file_metadata(b"%PDF-truncado", "roto.pdf")

    assert content_type == "application/pdf"
    assert _values(items)["Firma detectada"] == "PDF"
    assert any("PDF" in warning for warning in warnings)


def test_office_zip_bomb_is_not_decompressed(monkeypatch) -> None:
    _without_exiftool(monkeypatch)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/worksheets/sheet1.xml", "0" * 400_000)

    _content_type, _items, warnings = metadata.extract_file_metadata(
        output.getvalue(), "sospechoso.xlsx"
    )

    assert any("compresión sospechosa" in warning for warning in warnings)


def test_dwg_reports_only_header_level_coverage(monkeypatch) -> None:
    _without_exiftool(monkeypatch)

    content_type, items, warnings = metadata.extract_file_metadata(
        b"AC1032\x00contenido", "plano.dwg"
    )

    assert content_type == "application/acad"
    assert _values(items)["Versión de formato"] == "DWG 2018 (formato)"
    assert any("profunda" in warning for warning in warnings)


def test_exiftool_is_read_only_stdin_and_combined_with_native(monkeypatch) -> None:
    calls: list[tuple[list[str], dict]] = []

    def fake_popen(command, **kwargs):
        calls.append((command, kwargs))
        return _FakeProcess(
            io.BytesIO(b'[{"SourceFile":"-","File:FileName":"-","EXIF:Artist":"Ada"}]')
        )

    monkeypatch.setattr(exiftool.shutil, "which", lambda _name: "/opt/exiftool")
    monkeypatch.setattr(exiftool.subprocess, "Popen", fake_popen)

    _content_type, items, warnings = metadata.extract_file_metadata(
        b"adjunto desconocido", "nombre-malicioso.jpg"
    )

    command, kwargs = calls[0]
    assert command == ["/opt/exiftool", "-j", "-G1", "-s", "-"]
    assert all("FileName=" not in argument for argument in command)
    assert kwargs["shell"] is False
    assert kwargs["stdin"] is exiftool.subprocess.PIPE
    assert kwargs["stdout"] is exiftool.subprocess.PIPE
    assert kwargs["stderr"] is exiftool.subprocess.DEVNULL
    assert _values(items)["EXIF:Artist"] == "Ada"
    assert not warnings


def test_exiftool_values_are_truncated_before_serialization(monkeypatch) -> None:
    huge_value = "x" * (settings.max_property_chars + 100)
    exiftool_output = ('[{"XMP:Description":"' + huge_value + '"}]').encode()

    monkeypatch.setattr(exiftool.shutil, "which", lambda _name: "/opt/exiftool")
    monkeypatch.setattr(
        exiftool.subprocess,
        "Popen",
        lambda _command, **_kwargs: _FakeProcess(io.BytesIO(exiftool_output)),
    )

    _content_type, items, warnings = metadata.extract_file_metadata(b"x", "datos.bin")

    value = _values(items)["XMP:Description"]
    assert len(value) > settings.max_property_chars
    assert "truncado" in value
    assert any("truncó" in warning for warning in warnings)


def test_exiftool_output_larger_than_budget_is_killed(monkeypatch) -> None:
    process = _FakeProcess(io.BytesIO(b"x" * (exiftool.EXIFTOOL_STDOUT_MAX_BYTES + 1)))
    monkeypatch.setattr(exiftool.shutil, "which", lambda _name: "/opt/exiftool")
    monkeypatch.setattr(exiftool.subprocess, "Popen", lambda *_args, **_kwargs: process)

    _content_type, _items, warnings = metadata.extract_file_metadata(b"archivo", "datos.bin")

    assert process.killed is True
    assert process.returncode is not None
    assert any("más de 1 MiB" in warning for warning in warnings)


def test_exiftool_timeout_keeps_native_metadata_and_leaves_no_process(monkeypatch) -> None:
    image = Image.new("RGB", (3, 2))
    output = io.BytesIO()
    image.save(output, format="PNG")
    process = _FakeProcess(_BlockingStdout(), hangs=True)
    monkeypatch.setattr(exiftool.shutil, "which", lambda _name: "/opt/exiftool")
    monkeypatch.setattr(exiftool.subprocess, "Popen", lambda *_args, **_kwargs: process)

    _content_type, items, warnings = metadata.extract_file_metadata(
        output.getvalue(), "foto.png", external_deadline=time.monotonic() + 0.02
    )

    assert _values(items)["Dimensiones"] == "3 × 2"
    assert process.killed is True
    assert process.returncode is not None
    assert any("agotó el tiempo" in warning for warning in warnings)


def test_expired_external_deadline_skips_exiftool_but_keeps_native(monkeypatch) -> None:
    image = Image.new("RGB", (2, 1))
    output = io.BytesIO()
    image.save(output, format="PNG")
    monkeypatch.setattr(exiftool.shutil, "which", lambda _name: "/opt/exiftool")
    monkeypatch.setattr(
        exiftool.subprocess,
        "Popen",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("no debe iniciar ExifTool")),
    )

    _content_type, items, warnings = metadata.extract_file_metadata(
        output.getvalue(), "foto.png", external_deadline=time.monotonic() - 1
    )

    assert _values(items)["Dimensiones"] == "2 × 1"
    assert any("se omitió" in warning for warning in warnings)
